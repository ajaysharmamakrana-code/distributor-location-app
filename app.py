
from flask import Flask, render_template, jsonify
from openpyxl import load_workbook
import os

app = Flask(__name__)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MASTER_FILE = os.path.join(BASE_DIR, "Distributor_Master.xlsx")

def load_master():
    wb = load_workbook(MASTER_FILE, data_only=True, read_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    def norm(s):
        return " ".join(str(s).split()).lower()
    idx = {}
    for i, h in enumerate(headers):
        n = norm(h)
        if n == "so":
            idx["SO"] = i
        elif "distributor" in n and "name" in n:
            idx["Distributor Name"] = i
        elif "distributor" in n and "erp" in n:
            idx["Distributor Erp Id"] = i
        elif n.startswith("city"):
            idx["City"] = i
    required = ["SO","Distributor Name","Distributor Erp Id","City"]
    missing = [x for x in required if x not in idx]
    if missing:
        raise RuntimeError(f"Missing columns: {missing}")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        so = "" if r[idx["SO"]] is None else str(r[idx["SO"]]).strip()
        name = "" if r[idx["Distributor Name"]] is None else str(r[idx["Distributor Name"]]).strip()
        erp = "" if r[idx["Distributor Erp Id"]] is None else str(r[idx["Distributor Erp Id"]]).strip()
        city = "" if r[idx["City"]] is None else str(r[idx["City"]]).strip()
        if so and name and erp:
            rows.append({"SO":so,"Distributor Name":name,"Distributor Erp Id":erp,"City":city})
    return rows

@app.route("/")
def index():
    rows = load_master()
    sos = sorted({r["SO"] for r in rows})
    return render_template("index.html", sos=sos, master_rows=rows)

@app.route("/health")
def health():
    return jsonify({"ok": True})

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
